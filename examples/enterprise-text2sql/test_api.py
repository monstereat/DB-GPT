import json
import sqlite3
import sys
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from threading import Event, Thread
from types import SimpleNamespace

import api as api_module
import jwt
import pytest
from api import ReportMetricRequest, ReportRequest, _render_pdf_report, create_app
from cryptography.fernet import Fernet
from cryptography.hazmat.primitives.asymmetric import rsa
from ecommerce_demo import GOLD_QUERIES, create_demo_database
from fastapi import HTTPException
from fastapi.testclient import TestClient
from metric_governance import MetricReleaseStore
from openpyxl import load_workbook
from pypdf import PdfReader

from dbgpt_serve.utils import auth
from dbgpt_serve.utils.auth import UserRequest, get_user_from_headers


@pytest.fixture(autouse=True)
def configure_report_result_encryption(monkeypatch):
    monkeypatch.setenv(
        "DBGPT_REPORT_ENCRYPTION_KEY", Fernet.generate_key().decode("ascii")
    )
    monkeypatch.delenv("DBGPT_REPORT_ENCRYPTION_KEYS", raising=False)


def test_insecure_oidc_flag_only_allows_loopback_issuer(monkeypatch):
    monkeypatch.setenv("DBGPT_OIDC_AUDIENCE", "enterprise-text2sql-demo")
    monkeypatch.setenv("DBGPT_OIDC_ALLOW_INSECURE_HTTP", "true")
    for issuer in (
        "http://127.0.0.1:8080/realms/analytics",
        "http://[::1]:8080/realms/analytics",
        "http://localhost:8080/realms/analytics",
    ):
        monkeypatch.setenv("DBGPT_OIDC_ISSUER", issuer)
        assert auth._oidc_settings()["issuer"] == issuer

    for issuer in (
        "http://id.example.com/realms/analytics",
        "http://192.0.2.10/realms/analytics",
    ):
        monkeypatch.setenv("DBGPT_OIDC_ISSUER", issuer)
        with pytest.raises(RuntimeError, match="loopback HTTP"):
            auth._oidc_settings()


def test_insecure_oidc_flag_only_allows_loopback_jwks(monkeypatch):
    issuer = "http://127.0.0.1:8080/realms/analytics"

    class _DiscoveryResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return json.dumps(
                {"issuer": issuer, "jwks_uri": "http://id.example.com/jwks"}
            ).encode()

    monkeypatch.setenv("DBGPT_OIDC_ALLOW_INSECURE_HTTP", "true")
    monkeypatch.setattr(auth, "urlopen", lambda *_args, **_kwargs: _DiscoveryResponse())
    auth._get_jwks_client.cache_clear()
    with pytest.raises(RuntimeError, match="invalid jwks_uri"):
        auth._get_jwks_client(issuer)
    auth._get_jwks_client.cache_clear()


def test_oidc_settings_reject_empty_claim_paths_and_role_names(monkeypatch):
    monkeypatch.setenv("DBGPT_OIDC_ISSUER", "https://id.example.com/realms/analytics")
    monkeypatch.setenv("DBGPT_OIDC_AUDIENCE", "enterprise-text2sql-demo")
    monkeypatch.setenv("DBGPT_OIDC_ALLOW_INSECURE_HTTP", "false")

    for claim_mappings in (
        '{"tenant_id":""}',
        '{"tenant_id":"organization..tenant"}',
        '{"tenant_id":"organization. .tenant"}',
    ):
        monkeypatch.setenv("DBGPT_OIDC_CLAIM_MAPPINGS", claim_mappings)
        with pytest.raises(RuntimeError, match="non-empty claim paths"):
            auth._oidc_settings()

    monkeypatch.setenv("DBGPT_OIDC_CLAIM_MAPPINGS", "{}")
    for role_mapping in ('{"":"normal"}', '{"regional-sales":" "}'):
        monkeypatch.setenv("DBGPT_OIDC_ROLE_MAPPING", role_mapping)
        with pytest.raises(RuntimeError, match="non-empty role names"):
            auth._oidc_settings()


def test_query_uses_authenticated_tenant_and_ignores_body_tenant(tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-1", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/query",
            json={"sql": "SELECT product_id FROM products ORDER BY product_id"},
        )
        body_tenant = client.post(
            "/query",
            json={
                "sql": "SELECT product_id FROM products ORDER BY product_id",
                "tenant_id": "tenant-b",
                "region_id": "b-gz",
            },
        )

    assert response.status_code == 200
    assert response.json()["rows"] == [["a-p1"], ["a-p2"], ["a-p3"], ["a-p4"]]
    assert body_tenant.status_code == 422
    assert '"actor_user_id": "analyst-1"' in caplog.text
    assert '"policy_version": "1.0.0"' in caplog.text
    assert '"query_sha256":' in caplog.text
    assert "SELECT product_id" not in caplog.text


def test_query_fails_closed_without_a_mapped_or_allowed_tenant(tmp_path, caplog):
    caplog.set_level("WARNING")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-1"
    )

    with TestClient(app) as client:
        missing_tenant = client.post("/query", json={"sql": "SELECT 1"})
        app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
            user_id="analyst-2", tenant_id="tenant-unknown", role="normal"
        )
        unknown_tenant = client.post("/query", json={"sql": "SELECT 1"})

    assert missing_tenant.status_code == 403
    assert unknown_tenant.status_code == 403
    assert '"reason": "missing_tenant"' in caplog.text
    assert '"reason": "tenant_not_allowed"' in caplog.text
    assert "SELECT 1" not in caplog.text


def test_query_scopes_rows_from_verified_bearer_claim(tmp_path, monkeypatch):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    issuer = "https://id.example.test/realms/analytics"
    audience = "enterprise-text2sql-demo"
    private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)

    class _SigningKey:
        key = private_key.public_key()

    class _FakeJwksClient:
        def get_signing_key_from_jwt(self, token):
            return _SigningKey()

    monkeypatch.setattr(auth, "_get_jwks_client", lambda _: _FakeJwksClient())
    monkeypatch.setenv("DBGPT_OIDC_ISSUER", issuer)
    monkeypatch.setenv("DBGPT_OIDC_AUDIENCE", audience)
    monkeypatch.setenv(
        "DBGPT_OIDC_CLAIM_MAPPINGS", '{"tenant_id":"organization.tenant"}'
    )
    monkeypatch.setenv("DBGPT_OIDC_ROLE_MAPPING", '{"analytics-reader":"normal"}')
    token = jwt.encode(
        {
            "iss": issuer,
            "aud": audience,
            "sub": "analyst-a",
            "iat": int(time.time()),
            "exp": int(time.time()) + 60,
            "organization": {"tenant": "tenant-a"},
            "role": "analytics-reader",
        },
        private_key,
        algorithm="RS256",
    )

    with TestClient(app) as client:
        response = client.post(
            "/query",
            headers={"Authorization": f"Bearer {token}"},
            json={"sql": "SELECT product_id FROM products ORDER BY product_id"},
        )

    assert response.status_code == 200
    assert response.json()["rows"] == [["a-p1"], ["a-p2"], ["a-p3"], ["a-p4"]]


def test_oidc_region_claim_and_role_mapping_scope_sales_queries(tmp_path, monkeypatch):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    issuer = "https://id.example.test/realms/analytics"
    audience = "enterprise-text2sql-demo"
    private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)

    class _SigningKey:
        key = private_key.public_key()

    class _FakeJwksClient:
        def get_signing_key_from_jwt(self, token):
            return _SigningKey()

    monkeypatch.setattr(auth, "_get_jwks_client", lambda _: _FakeJwksClient())
    monkeypatch.setenv("DBGPT_OIDC_ISSUER", issuer)
    monkeypatch.setenv("DBGPT_OIDC_AUDIENCE", audience)
    monkeypatch.setenv(
        "DBGPT_OIDC_CLAIM_MAPPINGS",
        '{"tenant_id":"organization.tenant",'
        '"region_id":"organization.region","role":"realm_access.roles"}',
    )
    monkeypatch.setenv("DBGPT_OIDC_ROLE_MAPPING", '{"regional-sales":"sales"}')
    token = jwt.encode(
        {
            "iss": issuer,
            "aud": audience,
            "sub": "sales-a",
            "iat": int(time.time()),
            "exp": int(time.time()) + 60,
            "organization": {"tenant": "tenant-a", "region": "a-gz"},
            "realm_access": {"roles": ["regional-sales"]},
        },
        private_key,
        algorithm="RS256",
    )

    with TestClient(app) as client:
        response = client.post(
            "/query",
            headers={"Authorization": f"Bearer {token}"},
            json={"sql": "SELECT order_id FROM orders ORDER BY order_id"},
        )
        regional_metrics = client.post(
            "/query",
            headers={"Authorization": f"Bearer {token}"},
            json={"sql": GOLD_QUERIES["q41"]},
        )

    assert response.status_code == 200
    assert response.json()["rows"] == [[1], [2], [4], [8]]
    assert regional_metrics.status_code == 200
    assert regional_metrics.json()["rows"] == [[45000, 5.56]]


def test_oidc_uses_discovery_jwks_and_verified_rsa_token(tmp_path, monkeypatch):
    private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    jwk_data = jwt.algorithms.RSAAlgorithm.to_jwk(private_key.public_key())
    public_jwk = json.loads(jwk_data)
    public_jwk.update({"kid": "local-test-key", "use": "sig", "alg": "RS256"})

    class _OidcHandler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == "/realms/analytics/.well-known/openid-configuration":
                payload = {
                    "issuer": issuer,
                    "jwks_uri": f"{issuer}/protocol/openid-connect/certs",
                }
            elif self.path == "/realms/analytics/protocol/openid-connect/certs":
                payload = {"keys": [public_jwk]}
            else:
                self.send_error(404)
                return
            body = json.dumps(payload).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), _OidcHandler)
    issuer = f"http://127.0.0.1:{server.server_port}/realms/analytics"
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        auth._get_jwks_client.cache_clear()
        monkeypatch.setenv("DBGPT_OIDC_ISSUER", issuer)
        monkeypatch.setenv("DBGPT_OIDC_AUDIENCE", "enterprise-text2sql-demo")
        monkeypatch.setenv("DBGPT_OIDC_ALLOW_INSECURE_HTTP", "true")
        monkeypatch.setenv(
            "DBGPT_OIDC_CLAIM_MAPPINGS",
            '{"tenant_id":"organization.tenant",'
            '"region_id":"region_id","role":"realm_access.roles"}',
        )
        monkeypatch.setenv("DBGPT_OIDC_ROLE_MAPPING", '{"regional-sales":"sales"}')
        token = jwt.encode(
            {
                "iss": issuer,
                "aud": "enterprise-text2sql-demo",
                "sub": "sales-user-1",
                "iat": int(time.time()),
                "exp": int(time.time()) + 300,
                "organization": {"tenant": "tenant-a"},
                "region_id": "a-gz",
                "realm_access": {"roles": ["regional-sales"]},
            },
            private_key,
            algorithm="RS256",
            headers={"kid": "local-test-key"},
        )
        app = create_app(create_demo_database(tmp_path / "oidc-demo.sqlite"))

        with TestClient(app) as client:
            response = client.post(
                "/query",
                headers={"Authorization": f"Bearer {token}"},
                json={"sql": "SELECT order_id FROM orders ORDER BY order_id"},
            )

        assert response.status_code == 200, response.text
        assert response.json()["rows"] == [[1], [2], [4], [8]]
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)
        auth._get_jwks_client.cache_clear()


def test_metric_api_returns_registered_semantics_with_tenant_scope(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "sales_amount",
                "metric_version": "1.0.0",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["metric"]["metric_version"] == "1.0.0"
    assert body["result"]["rows"] == [["Guangzhou", 45000], ["Shenzhen", 32000]]


def test_available_metric_catalog_filters_admin_metrics_by_role(tmp_path):
    database = create_demo_database(tmp_path / "available-metrics.sqlite")
    app = create_app(database)
    identity = {"role": "normal"}
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="user-a", tenant_id="tenant-a", role=identity["role"]
    )

    with TestClient(app) as client:
        normal = client.get("/metrics/available")
        identity["role"] = "admin"
        admin = client.get("/metrics/available")

    assert normal.status_code == 200
    assert admin.status_code == 200
    normal_metrics = {item["id"]: item for item in normal.json()["metrics"]}
    admin_metrics = {item["id"]: item for item in admin.json()["metrics"]}
    assert "sales_amount" in normal_metrics
    assert "gross_margin_rate" not in normal_metrics
    assert admin_metrics["gross_margin_rate"] == {
        "id": "gross_margin_rate",
        "name": "毛利率",
        "unit": "percent",
        "versions": ["1.0.0"],
        "default_version": "1.0.0",
    }
    assert "internal_cost_cents" not in json.dumps(normal.json(), ensure_ascii=False)


def test_gross_margin_is_admin_only_and_raw_sql_still_denies_cost(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="admin-a", tenant_id="tenant-a", role="admin", region_id="a-gz"
    )

    with TestClient(app) as client:
        margin = client.post(
            "/metrics/query",
            json={
                "metric_id": "gross_margin_rate",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )
        raw_cost = client.post(
            "/query", json={"sql": "SELECT internal_cost_cents FROM products"}
        )

    assert margin.status_code == 200
    assert margin.json()["result"]["rows"] == [["Guangzhou", 48.0]]
    assert raw_cost.status_code == 403

    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    with TestClient(app) as client:
        denied = client.post(
            "/metrics/query",
            json={
                "metric_id": "gross_margin_rate",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )
    assert denied.status_code == 403


def test_report_export_uses_role_metric_and_field_allowlist(tmp_path, caplog):
    caplog.set_level("WARNING")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="admin-a", tenant_id="tenant-a", role="admin"
    )

    with TestClient(app) as client:
        denied_metric = client.post(
            "/reports/export",
            json={
                "metrics": [{"metric_id": "gross_margin_rate"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )
        denied_task = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "restricted-margin"},
            json={
                "metrics": [{"metric_id": "gross_margin_rate"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert denied_metric.status_code == 403
    assert denied_task.status_code == 403
    assert '"reason": "report_metric_not_exportable"' in caplog.text

    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="contractor-a", tenant_id="tenant-a", role="contractor"
    )
    with TestClient(app) as client:
        denied_role = client.post(
            "/reports/export",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )
    assert denied_role.status_code == 403


def test_report_templates_control_roles_metrics_and_export_columns(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    base_request = {
        "metrics": [{"metric_id": "sales_amount"}],
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
        "template_id": "sales-compact@1.0.0",
    }
    with TestClient(app) as client:
        compact = client.post("/reports/export", json=base_request)
        compact_pdf = client.post("/reports/export.pdf", json=base_request)
        denied_template = client.post(
            "/reports/export",
            json={**base_request, "template_id": "admin-margin@1.0.0"},
        )

    assert compact.status_code == 200
    workbook = load_workbook(BytesIO(compact.content), read_only=True, data_only=True)
    assert list(workbook["Metric Definitions"].values)[0] == (
        "metric_id",
        "metric_version",
        "definition",
    )
    assert list(workbook["Results"].values)[0] == (
        "metric_id",
        "dimension_value",
        "value",
    )
    assert denied_template.status_code == 403
    assert compact_pdf.status_code == 200
    assert compact_pdf.headers["content-type"] == "application/pdf"

    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="admin-a", tenant_id="tenant-a", role="admin"
    )
    with TestClient(app) as client:
        margin = client.post(
            "/reports/export",
            json={
                **base_request,
                "template_id": "admin-margin@1.0.0",
                "metrics": [{"metric_id": "gross_margin_rate"}],
            },
        )
    assert margin.status_code == 200
    workbook = load_workbook(BytesIO(margin.content), read_only=True, data_only=True)
    assert any(
        row[0] == "gross_margin_rate" for row in list(workbook["Results"].values)[1:]
    )


@pytest.mark.parametrize(
    ("path", "payload"),
    [
        (
            "/reports/export",
            {
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "fields": ["internal_cost_cents"],
            },
        ),
        (
            "/reports/export.pdf",
            {
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "fields": ["anything"],
            },
        ),
        (
            "/reports/tasks",
            {
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "fields": ["anything"],
            },
        ),
        (
            "/reports/batches/tasks",
            {
                "scenarios": [
                    {
                        "name": "safe scenario",
                        "report": {
                            "metrics": [{"metric_id": "sales_amount"}],
                            "start_date": "2026-04-01",
                            "end_date": "2026-07-01",
                            "fields": ["anything"],
                        },
                    },
                    {
                        "name": "second scenario",
                        "report": {
                            "metrics": [{"metric_id": "refund_rate"}],
                            "start_date": "2026-04-01",
                            "end_date": "2026-07-01",
                        },
                    },
                ]
            },
        ),
    ],
)
def test_report_apis_reject_client_template_fields_before_query_or_enqueue(
    tmp_path, monkeypatch, path, payload
):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database, report_task_db_path=tmp_path / "report-tasks.sqlite")
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    monkeypatch.setattr(
        api_module,
        "_collect_report_data",
        lambda *_args, **_kwargs: pytest.fail("invalid fields reached report query"),
    )
    monkeypatch.setattr(
        app.state.report_task_store,
        "enqueue",
        lambda *_args, **_kwargs: pytest.fail("invalid fields reached task enqueue"),
    )

    headers = {"Idempotency-Key": "reject-client-template"}
    with TestClient(app) as client:
        response = client.post(path, headers=headers, json=payload)

    assert response.status_code == 422


def test_unknown_report_template_is_denied_before_query_or_enqueue(
    tmp_path, monkeypatch
):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database, report_task_db_path=tmp_path / "report-tasks.sqlite")
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    monkeypatch.setattr(
        api_module,
        "_collect_report_data",
        lambda *_args, **_kwargs: pytest.fail("unknown template reached report query"),
    )
    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "template_id": "client-template",
            },
        )
    assert response.status_code == 403


def test_metric_api_allows_explicit_order_cohort_refund_version(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "paid_refund_amount",
                "metric_version": "2.0.0",
                "start_date": "2026-05-09",
                "end_date": "2026-05-11",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["metric"]["metric_version"] == "2.0.0"
    assert "原订单下单日期" in body["metric"]["definition"]
    assert body["result"]["rows"] == [[None]]


def test_metric_api_allows_cohort_net_sales_version(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    with sqlite3.connect(database) as connection:
        connection.execute(
            "INSERT INTO refunds VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                "tenant-a",
                99,
                1,
                "2026-07-03",
                1000,
                "paid",
                "synthetic cohort boundary",
            ),
        )
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "net_sales_amount",
                "metric_version": "2.0.0",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["metric"]["metric_version"] == "2.0.0"
    assert "cohort" in body["metric"]["definition"]
    assert body["result"]["rows"] == [[71300]]


def test_metric_api_allows_cohort_refund_rate_version(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    with sqlite3.connect(database) as connection:
        connection.execute(
            "INSERT INTO refunds VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                "tenant-a",
                99,
                1,
                "2026-07-03",
                1000,
                "paid",
                "synthetic cohort boundary",
            ),
        )
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "refund_rate",
                "metric_version": "2.0.0",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["metric"]["metric_version"] == "2.0.0"
    assert "cohort" in body["metric"]["definition"]
    assert body["result"]["rows"] == [[7.4]]


def test_metric_api_rejects_unpublished_version(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "sales_amount",
                "metric_version": "0.9.0",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 400


def test_xlsx_report_contains_scope_time_and_metric_definitions(tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [
                    {"metric_id": "sales_amount", "metric_version": "1.0.0"},
                    {"metric_id": "refund_rate"},
                ],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Report", "Metric Definitions", "Results"]
    overview = list(workbook["Report"].values)
    definitions = list(workbook["Metric Definitions"].values)
    results = list(workbook["Results"].values)
    assert overview[1][0] == "生成时间（UTC）"
    assert overview[2][:2] == ("租户", "tenant-a")
    assert definitions[1][0:2] == ("sales_amount", "1.0.0")
    assert definitions[0] == (
        "metric_id",
        "metric_version",
        "unit",
        "definition",
        "start_date",
        "end_date_exclusive",
        "dimension",
    )
    assert results[0] == (
        "metric_id",
        "metric_version",
        "dimension_value",
        "value",
        "unit",
    )
    assert any(row[2:4] == ("Guangzhou", 450) for row in results[1:])
    assert all(row[2] != "tenant-b" for row in results[1:])
    assert '"status": "report_exported"' in caplog.text


def test_xlsx_report_writes_formula_like_dimension_as_text(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE regions SET name = ? WHERE region_id = ?",
            ("=1+1", "a-gz"),
        )
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )

    assert response.status_code == 200
    workbook = load_workbook(
        BytesIO(response.content), read_only=False, data_only=False
    )
    cells = [
        cell
        for row in workbook["Results"].iter_rows()
        for cell in row
        if cell.value == "=1+1"
    ]
    assert len(cells) == 1
    assert cells[0].data_type == "s"


def test_export_validation_failure_is_audited_without_query_text(tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [
                    {"metric_id": "sales_amount"},
                    {"metric_id": "sales_amount"},
                ],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 400
    assert '"status": "report_export_failed"' in caplog.text
    assert '"format": "xlsx"' in caplog.text
    assert '"reason": "request_rejected"' in caplog.text
    assert "SELECT" not in caplog.text


def test_pdf_render_failure_is_audited(monkeypatch, tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    def fail_render(*_args):
        raise HTTPException(status_code=503, detail="PDF rendering unavailable")

    monkeypatch.setattr("api._render_pdf_report", fail_render)
    with TestClient(app) as client:
        response = client.post(
            "/reports/export.pdf",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 503
    assert '"status": "report_export_failed"' in caplog.text
    assert '"format": "pdf"' in caplog.text
    assert '"reason": "export_unavailable"' in caplog.text


def test_pdf_report_contains_scoped_metrics_and_is_audited(tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="sales-a",
        tenant_id="tenant-a",
        region_id="a-gz",
        role="sales",
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export.pdf",
            json={
                "metrics": [
                    {"metric_id": "sales_amount", "metric_version": "1.0.0"},
                    {"metric_id": "paid_refund_amount", "metric_version": "2.0.0"},
                ],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages
    )
    assert "tenant-a" in text
    assert "指标口径" in text
    assert "查询结果" in text
    assert "Guangzhou" in text
    assert "Shenzhen" not in text
    assert "2.0.0" in text
    assert "ORDER" not in text
    assert '"format": "pdf"' in caplog.text
    assert '"actor_user_id": "sales-a"' in caplog.text


def test_pdf_report_rejects_an_unavailable_configured_font(tmp_path, monkeypatch):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    monkeypatch.setenv("DBGPT_PDF_FONT", str(tmp_path / "missing-font.ttf"))

    with TestClient(app) as client:
        response = client.post(
            "/reports/export.pdf",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 503
    assert response.json()["detail"] == "Configured PDF font is unavailable"


def test_pdf_report_repeats_result_headers_across_pages():
    request = ReportRequest(
        metrics=[ReportMetricRequest(metric_id="sales_amount")],
        start_date="2026-04-01",
        end_date="2026-07-01",
        dimension="region",
    )
    user = UserRequest(user_id="analyst-a", tenant_id="tenant-a")
    definitions = [
        [
            "sales_amount",
            "1.0.0",
            "CNY_cent",
            "按订单创建日期统计的已完成销售额",
            "2026-04-01",
            "2026-07-01",
            "region",
        ]
    ]
    rows = [
        ["sales_amount", "1.0.0", f"区域 {index}", index * 100, "CNY"]
        for index in range(100)
    ]

    pdf = _render_pdf_report(
        request, user, "2026-09-25T10:00:00+00:00", definitions, rows
    )
    pages = PdfReader(BytesIO(pdf.getvalue())).pages

    assert len(pages) > 1
    page_text = [page.extract_text() or "" for page in pages]
    assert all("维度值" in text for text in page_text)
    assert "查询结果" in page_text[0]


def test_xlsx_report_applies_sales_region_scope(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="sales-a",
        tenant_id="tenant-a",
        region_id="a-gz",
        role="sales",
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )

    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    results = list(workbook["Results"].values)
    assert response.status_code == 200
    assert results[1:] == [("sales_amount", "1.0.0", "Guangzhou", 450, "CNY")]


def test_xlsx_report_rejects_duplicate_metrics(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/export",
            json={
                "metrics": [
                    {"metric_id": "sales_amount"},
                    {"metric_id": "sales_amount"},
                ],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 400


def test_async_report_task_is_durable_idempotent_and_scope_checked(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    task_database = tmp_path / "report-tasks.sqlite"
    app = create_app(database, report_task_db_path=task_database)
    current_user = {
        "value": UserRequest(user_id="analyst-a", tenant_id="tenant-a", role="normal")
    }
    app.dependency_overrides[get_user_from_headers] = lambda: current_user["value"]
    request = {
        "metrics": [{"metric_id": "sales_amount"}],
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
        "dimension": "region",
        "template_id": "sales-compact@1.0.0",
    }

    with TestClient(app) as client:
        created = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "sales-april"},
            json=request,
        )
        repeated = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "sales-april"},
            json=request,
        )
        mismatch = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "sales-april"},
            json={**request, "dimension": None},
        )

        assert created.status_code == 202
        assert repeated.status_code == 200
        assert repeated.json()["task_id"] == created.json()["task_id"]
        assert mismatch.status_code == 409

        task_id = created.json()["task_id"]
        deadline = time.monotonic() + 5
        task = None
        while time.monotonic() < deadline:
            response = client.get(f"/reports/tasks/{task_id}")
            assert response.status_code == 200
            task = response.json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "succeeded"
        assert task["attempts"] == 1
        assert task["download_url"] == f"/reports/tasks/{task_id}/export.xlsx"
        assert task["metric_catalog_sha256"]
        assert task["report_template_id"] == "sales-compact@1.0.0"
        assert task["report_template_version"] == "1.0.0"
        assert task["report_template_sha256"]

        download = client.get(task["download_url"])
        assert download.status_code == 200
        workbook = load_workbook(
            BytesIO(download.content), read_only=True, data_only=True
        )
        assert list(workbook["Report"].values)[2][:2] == ("租户", "tenant-a")
        assert any(
            row[0] == "sales_amount" and row[1] == "Guangzhou" and row[2] == 450
            for row in list(workbook["Results"].values)[1:]
        )
        assert list(workbook["Results"].values)[0] == (
            "metric_id",
            "dimension_value",
            "value",
        )

        current_user["value"] = UserRequest(
            user_id="analyst-a", tenant_id="tenant-b", role="normal"
        )
        assert client.get(f"/reports/tasks/{task_id}").status_code == 200
        assert client.get(task["download_url"]).status_code == 404
        current_user["value"] = UserRequest(
            user_id="analyst-b", tenant_id="tenant-a", role="normal"
        )
        assert client.get(f"/reports/tasks/{task_id}").status_code == 404

    # Recreating the app against the same task database preserves the result and key.
    restarted = create_app(database, report_task_db_path=task_database)
    restarted.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    with TestClient(restarted) as client:
        repeated = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "sales-april"},
            json=request,
        )
        assert repeated.status_code == 200
        assert repeated.json()["task_id"] == task_id
        assert repeated.json()["status"] == "succeeded"


def test_async_report_task_requires_result_encryption_key(tmp_path, monkeypatch):
    monkeypatch.delenv("DBGPT_REPORT_ENCRYPTION_KEY", raising=False)
    monkeypatch.delenv("DBGPT_REPORT_ENCRYPTION_KEYS", raising=False)
    database = create_demo_database(tmp_path / "demo.sqlite")
    task_database = tmp_path / "report-tasks.sqlite"
    app = create_app(database, report_task_db_path=task_database)
    assert app.state.report_task_worker is None
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "sales-no-encryption-key"},
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )

    assert response.status_code == 503
    assert response.json()["detail"] == "Report result encryption is not configured"
    with sqlite3.connect(task_database) as connection:
        count = connection.execute("SELECT COUNT(*) FROM report_tasks").fetchone()[0]
    assert count == 0


def test_async_report_batch_is_all_or_nothing_and_exports_named_scenarios(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database, report_task_db_path=tmp_path / "report-tasks.sqlite")
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    request = {
        "scenarios": [
            {
                "name": "Q2 区域销售",
                "report": {
                    "metrics": [{"metric_id": "sales_amount"}],
                    "start_date": "2026-04-01",
                    "end_date": "2026-07-01",
                    "dimension": "region",
                    "template_id": "sales-compact@1.0.0",
                },
            },
            {
                "name": "Q2 月度退款率",
                "report": {
                    "metrics": [{"metric_id": "refund_rate"}],
                    "start_date": "2026-04-01",
                    "end_date": "2026-07-01",
                    "dimension": "month",
                    "template_id": "sales-compact@1.0.0",
                },
            },
        ]
    }

    with TestClient(app) as client:
        created = client.post(
            "/reports/batches/tasks",
            headers={"Idempotency-Key": "q2-sales-and-refunds"},
            json=request,
        )
        task_id = created.json()["task_id"]
        deadline = time.monotonic() + 5
        task = None
        while time.monotonic() < deadline:
            task = client.get(f"/reports/tasks/{task_id}").json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)

        assert created.status_code == 202
        assert task["status"] == "succeeded"
        report = client.get(task["download_url"])
        assert report.status_code == 200
        workbook = load_workbook(
            BytesIO(report.content), read_only=True, data_only=True
        )
        definitions = list(workbook["Metric Definitions"].values)
        results = list(workbook["Results"].values)
        assert definitions[0] == (
            "scenario_name",
            "metric_id",
            "metric_version",
            "definition",
        )
        assert results[0] == (
            "scenario_name",
            "metric_id",
            "dimension_value",
            "value",
        )
        assert definitions[0][0] == "scenario_name"
        assert {row[0] for row in definitions[1:]} == {
            "Q2 区域销售",
            "Q2 月度退款率",
        }
        assert {row[0] for row in results[1:]} == {
            "Q2 区域销售",
            "Q2 月度退款率",
        }
        assert workbook["Report"]["B6"].value == "任一场景失败则整批失败"


def test_async_report_batch_validation_and_runtime_failure_are_atomic(
    tmp_path, monkeypatch
):
    database = create_demo_database(tmp_path / "demo.sqlite")
    task_database = tmp_path / "report-tasks.sqlite"
    app = create_app(database, report_task_db_path=task_database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    request = {
        "scenarios": [
            {
                "name": "first",
                "report": {
                    "metrics": [{"metric_id": "sales_amount"}],
                    "start_date": "2026-04-01",
                    "end_date": "2026-07-01",
                },
            },
            {
                "name": "second",
                "report": {
                    "metrics": [{"metric_id": "refund_rate"}],
                    "start_date": "2026-04-01",
                    "end_date": "2026-07-01",
                },
            },
        ]
    }
    original_collect = api_module._collect_report_data
    scenario_calls = 0

    def fail_on_second_scenario(*args, **kwargs):
        nonlocal scenario_calls
        scenario_calls += 1
        if scenario_calls % 2 == 0:
            raise RuntimeError("scenario failed")
        return original_collect(*args, **kwargs)

    monkeypatch.setattr(api_module, "_collect_report_data", fail_on_second_scenario)
    with TestClient(app) as client:
        too_small = client.post(
            "/reports/batches/tasks",
            headers={"Idempotency-Key": "one-scenario"},
            json={"scenarios": request["scenarios"][:1]},
        )
        duplicate_names = client.post(
            "/reports/batches/tasks",
            headers={"Idempotency-Key": "duplicate-scenario-names"},
            json={
                "scenarios": [
                    request["scenarios"][0],
                    {**request["scenarios"][1], "name": " FIRST "},
                ]
            },
        )
        created = client.post(
            "/reports/batches/tasks",
            headers={"Idempotency-Key": "atomic-failure"},
            json=request,
        )
        task_id = created.json()["task_id"]
        deadline = time.monotonic() + 5
        task = None
        while time.monotonic() < deadline:
            task = client.get(f"/reports/tasks/{task_id}").json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)

        assert too_small.status_code == 422
        assert duplicate_names.status_code == 422
        assert created.status_code == 202
        assert task["status"] == "failed"
        assert task["has_result"] is False
        assert task["download_url"] is None
        with sqlite3.connect(task_database) as connection:
            result_json = connection.execute(
                "SELECT result_json FROM report_tasks WHERE task_id = ?", (task_id,)
            ).fetchone()[0]
        assert result_json is None


def test_async_report_api_uses_active_key_from_configured_key_ring(
    tmp_path, monkeypatch
):
    active_key = Fernet.generate_key()
    old_key = Fernet.generate_key()
    monkeypatch.delenv("DBGPT_REPORT_ENCRYPTION_KEY", raising=False)
    monkeypatch.setenv(
        "DBGPT_REPORT_ENCRYPTION_KEYS",
        f"{active_key.decode('ascii')},{old_key.decode('ascii')}",
    )
    database = create_demo_database(tmp_path / "demo.sqlite")
    task_database = tmp_path / "encrypted-report-tasks.sqlite"
    app = create_app(database, report_task_db_path=task_database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        created = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "encrypted-sales"},
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "region",
            },
        )
        assert created.status_code == 202
        task_id = created.json()["task_id"]
        deadline = time.monotonic() + 5
        task = None
        while time.monotonic() < deadline:
            task = client.get(f"/reports/tasks/{task_id}").json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "succeeded"

        with sqlite3.connect(task_database) as connection:
            stored = connection.execute(
                "SELECT result_json FROM report_tasks WHERE task_id = ?",
                (task_id,),
            ).fetchone()[0]
        assert stored.startswith("enc:v1:")
        assert "Guangzhou" not in stored
        assert b"Guangzhou" in Fernet(active_key).decrypt(stored[7:].encode("ascii"))

        download = client.get(task["download_url"])
        assert download.status_code == 200
        workbook = load_workbook(
            BytesIO(download.content), read_only=True, data_only=True
        )
        assert any(
            row[0] == "sales_amount" and row[2] == "Guangzhou"
            for row in list(workbook["Results"].values)[1:]
        )


def test_async_report_requires_idempotency_key_and_owner(tmp_path):
    app = create_app(create_demo_database(tmp_path / "demo.sqlite"))
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        tenant_id="tenant-a", role="normal"
    )
    with TestClient(app) as client:
        missing_key = client.post(
            "/reports/tasks",
            json={
                "metrics": [{"metric_id": "sales_amount"}],
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
            },
        )
    assert missing_key.status_code == 403


def test_async_report_uses_catalog_snapshot_pinned_at_enqueue(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database, report_task_db_path=tmp_path / "report-tasks.sqlite")
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    worker = app.state.report_task_worker
    original_processor = worker.processor
    claimed = Event()
    resume = Event()

    def pause_before_processing(task):
        claimed.set()
        assert resume.wait(timeout=5)
        return original_processor(task)

    worker.processor = pause_before_processing
    request = {
        "metrics": [{"metric_id": "sales_amount"}],
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
    }
    with TestClient(app) as client:
        created = client.post(
            "/reports/tasks",
            headers={"Idempotency-Key": "pinned-catalog"},
            json=request,
        )
        assert created.status_code == 202
        assert claimed.wait(timeout=5)

        releases = MetricReleaseStore(database)
        catalog = releases.snapshot()
        next_version = next(
            item for item in catalog["metrics"] if item["id"] == "sales_amount"
        ).copy()
        next_version["version"] = "2.0.0"
        next_version["definition"] = "Published after this task was queued"
        releases.submit(
            next_version,
            actor_user_id="submitter-admin",
            actor_role="admin",
            tenant_id="tenant-a",
            reason="Snapshot regression test",
        )
        releases.review(
            "sales_amount",
            "2.0.0",
            approve=True,
            actor_user_id="reviewer-admin",
            actor_role="admin",
            tenant_id="tenant-a",
            reason="Snapshot regression test",
        )

        resume.set()
        task_id = created.json()["task_id"]
        deadline = time.monotonic() + 5
        task = None
        while time.monotonic() < deadline:
            task = client.get(f"/reports/tasks/{task_id}").json()
            if task["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.05)
        assert task["status"] == "succeeded"
        assert task["metric_catalog_version"] != releases.snapshot()["catalog_version"]

        report = client.get(task["download_url"])
        workbook = load_workbook(
            BytesIO(report.content), read_only=True, data_only=True
        )
        definitions = list(workbook["Metric Definitions"].values)
        assert definitions[1][1] == "1.0.0"


def test_metric_cache_rechecks_identity_and_audits_cache_hits(tmp_path, caplog):
    caplog.set_level("INFO")
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    current_user = {"value": UserRequest(user_id="analyst-a", tenant_id="tenant-a")}
    app.dependency_overrides[get_user_from_headers] = lambda: current_user["value"]
    request = {
        "metric_id": "sales_amount",
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
        "dimension": "region",
    }

    with TestClient(app) as client:
        first = client.post("/metrics/query", json=request)
        second = client.post("/metrics/query", json=request)
        current_user["value"] = UserRequest(user_id="analyst-b", tenant_id="tenant-a")
        other_user = client.post("/metrics/query", json=request)

    assert first.status_code == second.status_code == other_user.status_code == 200
    assert first.json()["cache"] == {"hit": False}
    assert second.json()["cache"] == {"hit": True}
    assert other_user.json()["cache"] == {"hit": False}
    assert '"status": "cache_hit"' in caplog.text
    assert '"actor_user_id": "analyst-a"' in caplog.text


def test_metric_cache_invalidates_when_database_changes(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )
    request = {
        "metric_id": "sales_amount",
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
        "dimension": "region",
    }

    with TestClient(app) as client:
        initial = client.post("/metrics/query", json=request)
        cached = client.post("/metrics/query", json=request)
        with sqlite3.connect(database) as connection:
            connection.execute(
                "UPDATE orders SET total_cents = total_cents + 1000 "
                "WHERE tenant_id = ? AND order_id = ?",
                ("tenant-a", 1),
            )
        updated = client.post("/metrics/query", json=request)

    assert initial.json()["cache"] == {"hit": False}
    assert cached.json()["cache"] == {"hit": True}
    assert updated.json()["cache"] == {"hit": False}
    assert updated.json()["result"]["rows"] == [
        ["Guangzhou", 46000],
        ["Shenzhen", 32000],
    ]


def test_configured_redis_cache_is_shared_between_api_instances(tmp_path, monkeypatch):
    database = create_demo_database(tmp_path / "demo.sqlite")
    shared_values = {}

    class FakeRedis:
        @classmethod
        def from_url(cls, url, **options):
            assert url == "redis://cache.example.test/0"
            assert options["decode_responses"] is True
            return cls()

        def get(self, key):
            return shared_values.get(key)

        def set(self, key, value, *, px):
            assert px > 0
            shared_values[key] = value

    monkeypatch.setitem(sys.modules, "redis", SimpleNamespace(Redis=FakeRedis))
    monkeypatch.setenv("DBGPT_METRIC_CACHE_REDIS_URL", "redis://cache.example.test/0")
    request = {
        "metric_id": "sales_amount",
        "start_date": "2026-04-01",
        "end_date": "2026-07-01",
        "dimension": "region",
    }
    first_app = create_app(database)
    second_app = create_app(database)
    other_user_app = create_app(database)
    first_app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a"
    )
    second_app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a"
    )
    other_user_app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-b", tenant_id="tenant-a"
    )

    with (
        TestClient(first_app) as first_client,
        TestClient(second_app) as second_client,
        TestClient(other_user_app) as other_user_client,
    ):
        first = first_client.post("/metrics/query", json=request)
        second = second_client.post("/metrics/query", json=request)
        with sqlite3.connect(database) as connection:
            connection.execute(
                "UPDATE orders SET total_cents = total_cents + 1000 "
                "WHERE tenant_id = ? AND order_id = ?",
                ("tenant-a", 1),
            )
        after_database_change = second_client.post("/metrics/query", json=request)
        other_user = other_user_client.post("/metrics/query", json=request)

    assert (
        first.status_code
        == second.status_code
        == after_database_change.status_code
        == 200
    )
    assert first.json()["cache"] == {"hit": False}
    assert second.json()["cache"] == {"hit": True}
    assert after_database_change.json()["cache"] == {"hit": False}
    assert after_database_change.json()["result"]["rows"] == [
        ["Guangzhou", 46000],
        ["Shenzhen", 32000],
    ]
    assert other_user.status_code == 200
    assert other_user.json()["cache"] == {"hit": False}


def test_metric_api_returns_monthly_trend_with_tenant_scope(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="analyst-a", tenant_id="tenant-a", role="normal"
    )

    with TestClient(app) as client:
        response = client.post(
            "/metrics/query",
            json={
                "metric_id": "sales_amount",
                "start_date": "2026-04-01",
                "end_date": "2026-07-01",
                "dimension": "month",
            },
        )
        previous_year = client.post(
            "/metrics/query",
            json={
                "metric_id": "sales_amount",
                "start_date": "2025-04-01",
                "end_date": "2025-07-01",
                "dimension": "month",
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["metric"]["dimension"] == "month"
    assert body["result"]["columns"] == ["period", "sales_cents"]
    assert body["result"]["rows"] == [
        ["2026-04", 10000],
        ["2026-05", 32000],
        ["2026-06", 35000],
    ]
    assert previous_year.status_code == 200
    assert previous_year.json()["result"]["rows"] == [["2025-04", 8000]]


def test_sales_api_enforces_region_claim_and_rejects_missing_region(tmp_path):
    database = create_demo_database(tmp_path / "demo.sqlite")
    app = create_app(database)
    app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
        user_id="sales-a",
        tenant_id="tenant-a",
        region_id="a-gz",
        role="sales",
    )

    with TestClient(app) as client:
        scoped = client.post("/query", json={"sql": "SELECT order_id FROM orders"})
        app.dependency_overrides[get_user_from_headers] = lambda: UserRequest(
            user_id="sales-no-region", tenant_id="tenant-a", role="sales"
        )
        missing_region = client.post(
            "/query", json={"sql": "SELECT order_id FROM orders"}
        )

    assert scoped.status_code == 200
    assert scoped.json()["rows"] == [[1], [2], [4], [8]]
    assert missing_region.status_code == 403
