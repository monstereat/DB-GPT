"""OIDC-authenticated API for the isolated multi-tenant SQLite demo."""

import hashlib
import json
import logging
import os
from datetime import datetime, timezone
from html import escape
from io import BytesIO
from pathlib import Path
from time import monotonic
from typing import Literal, Optional

from ecommerce_demo import (
    SCHEMA_METADATA,
    tenant_executor,
    tenant_gross_margin_executor,
)
from fastapi import Depends, FastAPI, Header, HTTPException, Response, status
from guarded_query import QueryRejected
from metric_cache import build_metric_cache, database_signature
from metric_governance import MetricReleaseStore
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, Field, field_validator
from report_tasks import (
    IdempotencyConflict,
    ReportResultEncryptionError,
    ReportTaskStore,
    ReportTaskWorker,
)
from report_template_governance import ReportTemplateReleaseStore
from semantic_metrics import build_metric_query
from starlette.responses import StreamingResponse

from dbgpt_serve.utils.auth import UserRequest, get_user_from_headers

logger = logging.getLogger(__name__)


def _catalog_snapshot_hash(snapshot: Optional[dict]) -> Optional[str]:
    if snapshot is None:
        return None
    payload = json.dumps(snapshot, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _resolve_report_template(
    template_id: str,
    role: str,
    snapshot: Optional[dict] = None,
    catalog: Optional[dict] = None,
) -> dict:
    catalog = catalog or REPORT_TEMPLATE_CATALOG
    template_by_id = {item["id"]: item for item in catalog["templates"]}
    template = (
        snapshot.get("template")
        if snapshot is not None
        else template_by_id.get(template_id)
    )
    if (
        not isinstance(template, dict)
        or template.get("id") != template_id
        or role not in template.get("allowed_roles", [])
        or not set(template.get("definition_fields", ()))
        <= set(REPORT_DEFINITION_VALUES)
        or not set(template.get("result_fields", ())) <= set(REPORT_RESULT_VALUES)
    ):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Report template is not available to this identity",
        )
    return template


def _project_report_rows(
    rows: list[list], fields: list[str], positions: dict
) -> list[list]:
    return [[row[positions[field]] for field in fields] for row in rows]


def _append_text_safe_row(sheet, values: list):
    """Append report values while keeping every string cell explicitly textual."""
    is_empty = (
        sheet.max_row == 1 and sheet.max_column == 1 and sheet["A1"].value is None
    )
    row_number = 1 if is_empty else sheet.max_row + 1
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_number, column=column, value=value)
        if isinstance(value, str):
            cell.data_type = "s"


class QueryRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sql: str


class MetricRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    metric_id: str
    start_date: str
    end_date: str
    dimension: Optional[Literal["region", "month"]] = None
    metric_version: Optional[str] = None


class MetricReleaseSubmission(BaseModel):
    model_config = ConfigDict(extra="forbid")

    definition: dict
    reason: str = Field(min_length=1, max_length=500)


class MetricReleaseReview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    approve: bool
    reason: str = Field(min_length=1, max_length=500)


class ReportTemplateReleaseSubmission(BaseModel):
    model_config = ConfigDict(extra="forbid")

    definition: dict
    reason: str = Field(min_length=1, max_length=500)


class ReportTemplateReleaseReview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    approve: bool
    reason: str = Field(min_length=1, max_length=500)


class ReportMetricRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    metric_id: str = Field(min_length=1, max_length=64, pattern=r"^[a-z][a-z0-9_]*$")
    metric_version: Optional[str] = None


class ReportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    metrics: list[ReportMetricRequest] = Field(min_length=1, max_length=4)
    start_date: str
    end_date: str
    dimension: Optional[Literal["region", "month"]] = None
    template_id: str = "sales-standard@1.0.0"


class ReportBatchScenario(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str = Field(min_length=1, max_length=80)
    report: ReportRequest

    @field_validator("name")
    @classmethod
    def scenario_name_must_not_be_blank(cls, name):
        name = name.strip()
        if not name:
            raise ValueError("Batch scenario name must not be blank")
        return name


class ReportBatchRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    scenarios: list[ReportBatchScenario] = Field(min_length=2, max_length=10)

    @field_validator("scenarios")
    @classmethod
    def scenario_names_must_be_unique(cls, scenarios):
        names = [scenario.name.strip().casefold() for scenario in scenarios]
        if any(not name for name in names) or len(set(names)) != len(names):
            raise ValueError("Batch scenario names must be non-empty and unique")
        return scenarios


REPORT_EXPORT_POLICY = {
    "sales_amount": frozenset({"admin", "normal", "sales"}),
    "paid_refund_amount": frozenset({"admin", "normal", "sales"}),
    "net_sales_amount": frozenset({"admin", "normal", "sales"}),
    "refund_rate": frozenset({"admin", "normal", "sales"}),
    "gross_margin_rate": frozenset({"admin"}),
}

REPORT_DEFINITION_FIELDS = (
    "metric_id",
    "metric_version",
    "unit",
    "definition",
    "start_date",
    "end_date_exclusive",
    "dimension",
)
REPORT_RESULT_FIELDS = (
    "metric_id",
    "metric_version",
    "dimension_value",
    "value",
    "unit",
)

REPORT_TEMPLATE_PATH = Path(__file__).with_name("report_templates.json")
with REPORT_TEMPLATE_PATH.open(encoding="utf-8") as template_file:
    REPORT_TEMPLATE_CATALOG = json.load(template_file)
REPORT_DEFINITION_VALUES = dict(zip(REPORT_DEFINITION_FIELDS, range(7)))
REPORT_RESULT_VALUES = dict(zip(REPORT_RESULT_FIELDS, range(5)))
REPORT_FIELD_LABELS = {
    "metric_id": "指标",
    "metric_version": "版本",
    "unit": "单位",
    "definition": "口径",
    "start_date": "开始日期",
    "end_date_exclusive": "结束日期（不含）",
    "dimension": "维度",
    "dimension_value": "维度值",
    "value": "结果",
}


def _audit_sink(user: UserRequest):
    def emit(event: dict):
        audit_event = {
            **event,
            "actor_user_id": user.user_id,
            "tenant_id": user.tenant_id,
            "region_id": user.region_id,
            "data_source_id": SCHEMA_METADATA["source"]["id"],
            "policy_version": SCHEMA_METADATA["source"]["policy_version"],
        }
        logger.info(
            "enterprise_query_audit %s",
            json.dumps(audit_event, sort_keys=True, ensure_ascii=False),
        )

    return emit


def _audit_authorization_rejection(user: UserRequest, reason: str):
    audit_event = {
        "actor_user_id": user.user_id,
        "tenant_id": user.tenant_id,
        "region_id": user.region_id,
        "data_source_id": SCHEMA_METADATA["source"]["id"],
        "policy_version": SCHEMA_METADATA["source"]["policy_version"],
        "status": "rejected",
        "reason": reason,
    }
    logger.warning(
        "enterprise_query_audit %s",
        json.dumps(audit_event, sort_keys=True, ensure_ascii=False),
    )


def _audit_export_failure(
    user: UserRequest, report_format: str, request: ReportRequest, status_code: int
):
    if status_code == status.HTTP_400_BAD_REQUEST:
        reason = "request_rejected"
    elif status_code == status.HTTP_403_FORBIDDEN:
        reason = "query_policy_rejected"
    elif status_code == status.HTTP_503_SERVICE_UNAVAILABLE:
        reason = "export_unavailable"
    else:
        reason = "export_failed"
    _audit_sink(user)(
        {
            "status": "report_export_failed",
            "format": report_format,
            "status_code": status_code,
            "reason": reason,
            "metric_count": len(request.metrics),
        }
    )


def _executor_for_user(database_path: str | Path, user: UserRequest):
    if not user.tenant_id:
        _audit_authorization_rejection(user, "missing_tenant")
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Authenticated identity has no tenant claim",
        )
    if user.role == "sales" and not user.region_id:
        _audit_authorization_rejection(user, "missing_region")
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Sales identity has no region claim",
        )
    try:
        return tenant_executor(
            database_path,
            user.tenant_id,
            region_id=user.region_id,
            audit_sink=_audit_sink(user),
        )
    except ValueError as exc:
        _audit_authorization_rejection(user, "tenant_not_allowed")
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Authenticated tenant is not authorized for this demo",
        ) from exc


def _validate_report_request(
    request: ReportRequest,
    user: UserRequest,
    metric_catalog: dict | None = None,
    template: dict | None = None,
) -> None:
    template = template or _resolve_report_template(request.template_id, user.role)
    if len({item.metric_id for item in request.metrics}) != len(request.metrics):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Report metric IDs must be unique",
        )
    for selection in request.metrics:
        if selection.metric_id not in template["allowed_metric_ids"]:
            _audit_authorization_rejection(user, "report_metric_not_exportable")
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Requested metric is not available for report export",
            )
        allowed_roles = REPORT_EXPORT_POLICY.get(selection.metric_id)
        if allowed_roles is None or user.role not in allowed_roles:
            _audit_authorization_rejection(user, "report_metric_not_exportable")
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Requested metric is not available for report export",
            )
        try:
            build_metric_query(
                selection.metric_id,
                request.start_date,
                request.end_date,
                dimension=request.dimension,
                metric_version=selection.metric_version,
                role=user.role,
                metric_catalog=metric_catalog,
            )
        except ValueError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Metric, version, or date range is not allowed",
            ) from exc


def _collect_report_data(
    database_path: str | Path,
    request: ReportRequest,
    user: UserRequest,
    metric_catalog: dict | None = None,
    template: dict | None = None,
):
    _validate_report_request(request, user, metric_catalog, template)
    executor = _executor_for_user(database_path, user)
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    definitions = []
    result_rows = []
    try:
        for selection in request.metrics:
            metric = build_metric_query(
                selection.metric_id,
                request.start_date,
                request.end_date,
                dimension=request.dimension,
                metric_version=selection.metric_version,
                role=user.role,
                metric_catalog=metric_catalog,
            )
            metric_executor = executor
            if selection.metric_id == "gross_margin_rate":
                metric_executor = tenant_gross_margin_executor(
                    database_path,
                    user.tenant_id,
                    role=user.role,
                    region_id=user.region_id,
                    audit_sink=_audit_sink(user),
                )
            result = metric_executor.run(metric["sql"])
            definitions.append(
                [
                    metric["metric_id"],
                    metric["metric_version"],
                    metric["unit"],
                    metric["definition"],
                    metric["start_date"],
                    metric["end_date"],
                    metric["dimension"] or "overall",
                ]
            )
            for row in result["rows"]:
                dimension_value = row[0] if request.dimension else "整体"
                value = row[-1]
                if metric["unit"] == "CNY_cent" and value is not None:
                    value = value / 100
                result_rows.append(
                    [
                        metric["metric_id"],
                        metric["metric_version"],
                        dimension_value,
                        value,
                        "CNY" if metric["unit"] == "CNY_cent" else metric["unit"],
                    ]
                )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Metric, version, or date range is not allowed",
        ) from exc
    except QueryRejected as exc:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Report query rejected by the demo query policy",
        ) from exc
    return generated_at, definitions, result_rows


def _render_xlsx_report(
    request: ReportRequest,
    user: UserRequest,
    generated_at: str,
    definitions: list[list],
    result_rows: list[list],
    template: dict | None = None,
) -> BytesIO:
    template = template or _resolve_report_template(request.template_id, user.role)
    definition_fields = template["definition_fields"]
    result_fields = template["result_fields"]
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Report"
    _append_text_safe_row(overview, ["DB-GPT 经营分析报告"])
    _append_text_safe_row(overview, ["生成时间（UTC）", generated_at])
    _append_text_safe_row(overview, ["租户", user.tenant_id])
    _append_text_safe_row(overview, ["区域范围", user.region_id or "全部授权区域"])
    _append_text_safe_row(
        overview, ["统计区间（左闭右开）", request.start_date, request.end_date]
    )
    _append_text_safe_row(overview, ["统计维度", request.dimension or "整体"])

    definitions_sheet = workbook.create_sheet("Metric Definitions")
    _append_text_safe_row(
        definitions_sheet,
        definition_fields,
    )
    for row in _project_report_rows(
        definitions, definition_fields, REPORT_DEFINITION_VALUES
    ):
        _append_text_safe_row(definitions_sheet, row)

    results_sheet = workbook.create_sheet("Results")
    _append_text_safe_row(
        results_sheet,
        result_fields,
    )
    for row in _project_report_rows(result_rows, result_fields, REPORT_RESULT_VALUES):
        _append_text_safe_row(results_sheet, row)
    for sheet in (overview, definitions_sheet, results_sheet):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 28
        for column in ("B", "C", "D", "E", "F", "G"):
            sheet.column_dimensions[column].width = 24

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _render_xlsx_batch_report(
    batch: ReportBatchRequest,
    user: UserRequest,
    generated_at: str,
    scenarios: list[dict],
    template: dict | None = None,
) -> BytesIO:
    template = template or _resolve_report_template(
        batch.scenarios[0].report.template_id, user.role
    )
    definition_fields = template["definition_fields"]
    result_fields = template["result_fields"]
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Report"
    _append_text_safe_row(overview, ["DB-GPT 批量经营分析报告"])
    _append_text_safe_row(overview, ["生成时间（UTC）", generated_at])
    _append_text_safe_row(overview, ["租户", user.tenant_id])
    _append_text_safe_row(overview, ["区域范围", user.region_id or "全部授权区域"])
    _append_text_safe_row(overview, ["场景数", len(batch.scenarios)])
    _append_text_safe_row(overview, ["失败策略", "任一场景失败则整批失败"])
    _append_text_safe_row(overview, ["场景名称"])
    for scenario in batch.scenarios:
        _append_text_safe_row(overview, [scenario.name])

    definitions_sheet = workbook.create_sheet("Metric Definitions")
    _append_text_safe_row(definitions_sheet, ["scenario_name", *definition_fields])
    results_sheet = workbook.create_sheet("Results")
    _append_text_safe_row(results_sheet, ["scenario_name", *result_fields])
    for scenario in scenarios:
        scenario_name = scenario["name"]
        for row in scenario["definitions"]:
            projected = _project_report_rows(
                [row], definition_fields, REPORT_DEFINITION_VALUES
            )[0]
            _append_text_safe_row(definitions_sheet, [scenario_name, *projected])
        for row in scenario["result_rows"]:
            projected = _project_report_rows(
                [row], result_fields, REPORT_RESULT_VALUES
            )[0]
            _append_text_safe_row(results_sheet, [scenario_name, *projected])

    for sheet in (overview, definitions_sheet, results_sheet):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column).column_letter
            ].width = 24

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _task_identity(user: UserRequest) -> dict:
    return {
        "user_id": user.user_id,
        "tenant_id": user.tenant_id,
        "region_id": user.region_id,
        "role": user.role,
        "data_source_id": SCHEMA_METADATA["source"]["id"],
        "authorization_policy_version": SCHEMA_METADATA["source"]["policy_version"],
    }


def _require_task_owner(user: UserRequest) -> str:
    if not user.user_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Authenticated user ID is required for report tasks",
        )
    return user.user_id


def _render_pdf_report(
    request: ReportRequest,
    user: UserRequest,
    generated_at: str,
    definitions: list[list],
    result_rows: list[list],
    template: dict | None = None,
) -> BytesIO:
    template = template or _resolve_report_template(request.template_id, user.role)
    definition_fields = template["definition_fields"]
    result_fields = template["result_fields"]
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            LongTable,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PDF export dependency is unavailable",
        ) from exc

    configured_font = os.getenv("DBGPT_PDF_FONT")
    if configured_font:
        try:
            pdfmetrics.registerFont(TTFont("DBGPT-CJK", configured_font))
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Configured PDF font is unavailable",
            ) from exc
        font_name = "DBGPT-CJK"
    else:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        font_name = "STSong-Light"
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=11,
    )
    title = ParagraphStyle(
        "ReportTitle", parent=body, fontSize=18, leading=24, alignment=TA_CENTER
    )
    section = ParagraphStyle(
        "ReportSection",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=12,
        leading=16,
    )
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
        title="DB-GPT 经营分析报告",
    )

    def paragraph(value):
        return Paragraph(escape(str(value if value is not None else "")), body)

    metadata = [
        [paragraph("生成时间（UTC）"), paragraph(generated_at)],
        [paragraph("租户"), paragraph(user.tenant_id)],
        [paragraph("区域范围"), paragraph(user.region_id or "全部授权区域")],
        [
            paragraph("统计区间（左闭右开）"),
            paragraph(f"{request.start_date} 至 {request.end_date}"),
        ],
        [paragraph("统计维度"), paragraph(request.dimension or "整体")],
    ]
    definitions_table = [
        [paragraph(REPORT_FIELD_LABELS[field]) for field in definition_fields],
        *[
            [paragraph(value) for value in row]
            for row in _project_report_rows(
                definitions, definition_fields, REPORT_DEFINITION_VALUES
            )
        ],
    ]
    results_table = [
        [paragraph(REPORT_FIELD_LABELS[field]) for field in result_fields],
        *[
            [paragraph(value) for value in row]
            for row in _project_report_rows(
                result_rows, result_fields, REPORT_RESULT_VALUES
            )
        ],
    ]
    table_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf1ef")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d5dfdc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )
    story = [
        Paragraph("DB-GPT 经营分析报告", title),
        Spacer(1, 14),
        Table(metadata, colWidths=[120, 385], style=table_style),
        Spacer(1, 14),
        Paragraph("指标口径", section),
        LongTable(definitions_table, repeatRows=1, style=table_style),
        Spacer(1, 14),
        Paragraph("查询结果", section),
        LongTable(results_table, repeatRows=1, style=table_style),
    ]
    document.build(story)
    buffer.seek(0)
    return buffer


def create_app(
    database_path: Optional[str | Path] = None,
    report_task_db_path: Optional[str | Path] = None,
    report_template_database_url: Optional[str] = None,
) -> FastAPI:
    app = FastAPI(title="Enterprise Text-to-SQL Demo")
    metric_cache = build_metric_cache(os.getenv("DBGPT_METRIC_CACHE_REDIS_URL"))
    release_store = MetricReleaseStore(database_path) if database_path else None
    template_database = (
        report_template_database_url
        or os.getenv("DBGPT_REPORT_TEMPLATE_DATABASE_URL")
        or database_path
    )
    template_store = (
        ReportTemplateReleaseStore(
            template_database,
            allowed_metrics=REPORT_EXPORT_POLICY,
            definition_fields=REPORT_DEFINITION_FIELDS,
            result_fields=REPORT_RESULT_FIELDS,
            seed_path=REPORT_TEMPLATE_PATH,
        )
        if template_database
        else None
    )

    def current_template_catalog() -> dict:
        return template_store.snapshot() if template_store else REPORT_TEMPLATE_CATALOG

    def resolve_template(
        template_id: str, role: str, snapshot: Optional[dict] = None
    ) -> dict:
        return _resolve_report_template(
            template_id,
            role,
            snapshot=snapshot,
            catalog=current_template_catalog(),
        )

    def capture_template_snapshot(template: dict) -> dict:
        return {
            "catalog_version": current_template_catalog()["catalog_version"],
            "template": json.loads(json.dumps(template)),
        }

    if report_task_db_path is None:
        configured_task_db = os.getenv("DBGPT_REPORT_TASK_DATABASE_URL") or os.getenv(
            "DBGPT_REPORT_TASK_DB"
        )
        if configured_task_db:
            report_task_db_path = configured_task_db
        elif database_path is not None:
            report_task_db_path = f"{database_path}.report-tasks.sqlite3"
    task_store = ReportTaskStore(report_task_db_path) if report_task_db_path else None
    app.state.report_template_store = template_store

    def process_report_task(task: dict) -> dict:
        if database_path is None:
            raise RuntimeError("demo database is not configured")
        if (
            task["identity"].get("data_source_id") != SCHEMA_METADATA["source"]["id"]
            or task["identity"].get("authorization_policy_version")
            != SCHEMA_METADATA["source"]["policy_version"]
        ):
            raise RuntimeError("authorization snapshot is stale")
        request_payload = task["request"]
        metric_catalog = request_payload.get("metric_catalog")
        template_snapshot = request_payload.get("report_template")
        template_hash = request_payload.get("report_template_sha256")
        if (
            template_snapshot is None
            or _catalog_snapshot_hash(template_snapshot) != template_hash
        ):
            raise RuntimeError("report template snapshot is invalid")
        user = UserRequest(
            user_id=task["identity"].get("user_id"),
            tenant_id=task["identity"].get("tenant_id"),
            region_id=task["identity"].get("region_id"),
            role=task["identity"].get("role"),
        )
        template_id = (
            ReportBatchRequest.model_validate(request_payload["report_batch"])
            .scenarios[0]
            .report.template_id
            if "report_batch" in request_payload
            else ReportRequest.model_validate(
                request_payload["report_request"]
            ).template_id
        )
        template = resolve_template(template_id, user.role, template_snapshot)
        if "report_batch" in request_payload:
            batch = ReportBatchRequest.model_validate(request_payload["report_batch"])
            scenarios = []
            for scenario in batch.scenarios:
                generated_at, definitions, result_rows = _collect_report_data(
                    database_path,
                    scenario.report,
                    user,
                    metric_catalog,
                    template,
                )
                scenarios.append(
                    {
                        "name": scenario.name,
                        "generated_at": generated_at,
                        "definitions": definitions,
                        "result_rows": result_rows,
                    }
                )
            return {
                "generated_at": datetime.now(timezone.utc).isoformat(
                    timespec="seconds"
                ),
                "scenarios": scenarios,
            }

        request = ReportRequest.model_validate(request_payload["report_request"])
        generated_at, definitions, result_rows = _collect_report_data(
            database_path, request, user, metric_catalog, template
        )
        return {
            "generated_at": generated_at,
            "definitions": definitions,
            "result_rows": result_rows,
        }

    worker = (
        ReportTaskWorker(task_store, process_report_task)
        if task_store and task_store.result_encryption_enabled
        else None
    )
    app.state.report_task_store = task_store
    app.state.report_task_worker = worker

    @app.on_event("startup")
    def start_report_worker():
        if worker:
            worker.start()

    @app.on_event("shutdown")
    def stop_report_worker():
        if worker:
            worker.stop()

    @app.post("/query")
    def query(
        request: QueryRequest,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if database_path is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Demo database is not configured",
            )
        executor = _executor_for_user(database_path, user)
        try:
            return executor.run(request.sql)
        except QueryRejected as exc:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="SQL rejected by the demo query policy",
            ) from exc

    @app.post("/metrics/query")
    def query_metric(
        request: MetricRequest,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if database_path is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Demo database is not configured",
            )
        if request.metric_id == "gross_margin_rate" and user.role != "admin":
            _audit_authorization_rejection(user, "restricted_metric_denied")
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Metric is not available to this identity",
            )
        try:
            metric = build_metric_query(
                request.metric_id,
                request.start_date,
                request.end_date,
                dimension=request.dimension,
                metric_version=request.metric_version,
                role=user.role,
                metric_catalog=release_store.snapshot() if release_store else None,
            )
        except ValueError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Metric or date range is not allowed",
            ) from exc
        executor = _executor_for_user(database_path, user)
        if request.metric_id == "gross_margin_rate":
            try:
                executor = tenant_gross_margin_executor(
                    database_path,
                    user.tenant_id,
                    role=user.role,
                    region_id=user.region_id,
                    audit_sink=_audit_sink(user),
                )
            except ValueError as exc:
                _audit_authorization_rejection(user, "restricted_metric_denied")
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Metric is not available to this identity",
                ) from exc
        sql_hash = hashlib.sha256(metric["sql"].encode("utf-8")).hexdigest()
        database_version = database_signature(database_path)
        cache_key = (
            SCHEMA_METADATA["source"]["id"],
            SCHEMA_METADATA["source"]["policy_version"],
            database_version,
            user.user_id,
            user.tenant_id,
            user.region_id,
            metric["metric_id"],
            metric["metric_version"],
            metric.get("content_sha256"),
            request.start_date,
            request.end_date,
            request.dimension,
            sql_hash,
        )
        started = monotonic()
        cached_result = metric_cache.get(cache_key)
        if (
            cached_result is not None
            and database_signature(database_path) == database_version
        ):
            duration_ms = int((monotonic() - started) * 1000)
            cached_result["duration_ms"] = duration_ms
            _audit_sink(user)(
                {
                    "query_sha256": sql_hash,
                    "metric_content_sha256": metric.get("content_sha256"),
                    "metric_id": metric["metric_id"],
                    "metric_version": metric["metric_version"],
                    "status": "cache_hit",
                    "duration_ms": duration_ms,
                    "returned_rows": cached_result["returned_rows"],
                    "allowed_tables": sorted(executor.allowed_tables),
                }
            )
            return {
                "metric": metric,
                "result": cached_result,
                "cache": {"hit": True},
            }
        try:
            result = executor.run(metric["sql"])
        except QueryRejected as exc:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Metric query rejected by the demo query policy",
            ) from exc
        _audit_sink(user)(
            {
                "query_sha256": sql_hash,
                "metric_content_sha256": metric.get("content_sha256"),
                "metric_id": metric["metric_id"],
                "metric_version": metric["metric_version"],
                "status": "success",
                "returned_rows": result["returned_rows"],
                "allowed_tables": sorted(executor.allowed_tables),
            }
        )
        if database_signature(database_path) == database_version:
            metric_cache.set(cache_key, result)
        return {"metric": metric, "result": result, "cache": {"hit": False}}

    @app.get("/metrics/available")
    def available_metrics(user: UserRequest = Depends(get_user_from_headers)):
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        snapshot = release_store.snapshot()
        visible_metrics = {}
        for metric in snapshot["metrics"]:
            allowed_roles = metric.get("allowed_roles")
            if allowed_roles and user.role not in allowed_roles:
                continue
            metric_id = metric.get("id")
            name = metric.get("name")
            unit = metric.get("unit")
            version = metric.get("version")
            if (
                not isinstance(metric_id, str)
                or not isinstance(name, str)
                or unit not in {"CNY_cent", "percent"}
                or not isinstance(version, str)
            ):
                continue
            visible = visible_metrics.setdefault(
                metric_id,
                {
                    "id": metric_id,
                    "name": name,
                    "unit": unit,
                    "versions": [],
                },
            )
            visible["versions"].append(version)
        defaults = snapshot["default_metric_versions"]
        metrics = list(visible_metrics.values())
        for metric in metrics:
            metric["default_version"] = defaults.get(metric["id"])
        return {"catalog_version": snapshot["catalog_version"], "metrics": metrics}

    @app.get("/metrics/catalog")
    def metric_catalog(user: UserRequest = Depends(get_user_from_headers)):
        if user.role != "admin":
            _audit_authorization_rejection(user, "metric_catalog_admin_required")
            raise HTTPException(status_code=403, detail="Admin role is required")
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        return release_store.snapshot()

    @app.get("/metrics/releases")
    def metric_releases(user: UserRequest = Depends(get_user_from_headers)):
        if user.role != "admin":
            _audit_authorization_rejection(user, "metric_catalog_admin_required")
            raise HTTPException(status_code=403, detail="Admin role is required")
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        return release_store.releases()

    @app.get("/metrics/releases/audit")
    def metric_release_audit(user: UserRequest = Depends(get_user_from_headers)):
        if user.role != "admin":
            _audit_authorization_rejection(user, "metric_catalog_admin_required")
            raise HTTPException(status_code=403, detail="Admin role is required")
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        return release_store.audit_events()

    @app.post("/metrics/releases")
    def submit_metric_release(
        request: MetricReleaseSubmission,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        try:
            return release_store.submit(
                request.definition,
                actor_user_id=user.user_id,
                actor_role=user.role,
                tenant_id=user.tenant_id,
                reason=request.reason,
            )
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.post("/metrics/releases/{metric_id}/{version}/review")
    def review_metric_release(
        metric_id: str,
        version: str,
        request: MetricReleaseReview,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if release_store is None:
            raise HTTPException(
                status_code=503, detail="Metric release storage is not configured"
            )
        try:
            return release_store.review(
                metric_id,
                version,
                approve=request.approve,
                actor_user_id=user.user_id,
                actor_role=user.role,
                tenant_id=user.tenant_id,
                reason=request.reason,
            )
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.get("/reports/templates/available")
    def available_report_templates(user: UserRequest = Depends(get_user_from_headers)):
        if template_store is None:
            raise HTTPException(
                status_code=503, detail="Report template storage is not configured"
            )
        snapshot = template_store.snapshot()
        templates = [
            template
            for template in snapshot["templates"]
            if user.role in template["allowed_roles"]
        ]
        return {"catalog_version": snapshot["catalog_version"], "templates": templates}

    @app.get("/reports/templates/releases")
    def report_template_releases(user: UserRequest = Depends(get_user_from_headers)):
        if user.role != "admin":
            raise HTTPException(status_code=403, detail="Admin role is required")
        if template_store is None:
            raise HTTPException(
                status_code=503, detail="Report template storage is not configured"
            )
        return template_store.releases()

    @app.get("/reports/templates/releases/audit")
    def report_template_audit(user: UserRequest = Depends(get_user_from_headers)):
        if user.role != "admin":
            raise HTTPException(status_code=403, detail="Admin role is required")
        if template_store is None:
            raise HTTPException(
                status_code=503, detail="Report template storage is not configured"
            )
        return template_store.audit_events()

    @app.post("/reports/templates/releases")
    def submit_report_template_release(
        request: ReportTemplateReleaseSubmission,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if template_store is None:
            raise HTTPException(
                status_code=503, detail="Report template storage is not configured"
            )
        try:
            return template_store.submit(
                request.definition,
                actor_user_id=user.user_id,
                actor_role=user.role,
                reason=request.reason,
            )
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.post("/reports/templates/releases/{template_id}/review")
    def review_report_template_release(
        template_id: str,
        request: ReportTemplateReleaseReview,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if template_store is None:
            raise HTTPException(
                status_code=503, detail="Report template storage is not configured"
            )
        try:
            return template_store.review(
                template_id,
                approve=request.approve,
                actor_user_id=user.user_id,
                actor_role=user.role,
                reason=request.reason,
            )
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.post("/reports/export")
    def export_report(
        request: ReportRequest,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if database_path is None:
            _audit_export_failure(
                user, "xlsx", request, status.HTTP_503_SERVICE_UNAVAILABLE
            )
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Demo database is not configured",
            )
        template = resolve_template(request.template_id, user.role)
        try:
            generated_at, definitions, result_rows = _collect_report_data(
                database_path,
                request,
                user,
                release_store.snapshot() if release_store else None,
                template,
            )
        except HTTPException as exc:
            _audit_export_failure(user, "xlsx", request, exc.status_code)
            raise

        buffer = _render_xlsx_report(
            request, user, generated_at, definitions, result_rows, template
        )
        _audit_sink(user)(
            {
                "status": "report_exported",
                "format": "xlsx",
                "metric_count": len(definitions),
                "returned_rows": len(result_rows),
                "generated_at": generated_at,
            }
        )
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": (
                    'attachment; filename="enterprise-analytics-report.xlsx"'
                )
            },
        )

    @app.post("/reports/export.pdf")
    def export_pdf_report(
        request: ReportRequest,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if database_path is None:
            _audit_export_failure(
                user, "pdf", request, status.HTTP_503_SERVICE_UNAVAILABLE
            )
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Demo database is not configured",
            )
        template = resolve_template(request.template_id, user.role)
        try:
            generated_at, definitions, result_rows = _collect_report_data(
                database_path,
                request,
                user,
                release_store.snapshot() if release_store else None,
                template,
            )
            buffer = _render_pdf_report(
                request, user, generated_at, definitions, result_rows, template
            )
        except HTTPException as exc:
            _audit_export_failure(user, "pdf", request, exc.status_code)
            raise
        _audit_sink(user)(
            {
                "status": "report_exported",
                "format": "pdf",
                "metric_count": len(definitions),
                "returned_rows": len(result_rows),
                "generated_at": generated_at,
            }
        )
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={
                "Content-Disposition": (
                    'attachment; filename="enterprise-analytics-report.pdf"'
                )
            },
        )

    @app.post("/reports/tasks")
    def create_report_task(
        request: ReportRequest,
        response: Response,
        idempotency_key: Optional[str] = Header(None, alias="Idempotency-Key"),
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if task_store is None or database_path is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report task storage or demo database is not configured",
            )
        actor_user_id = _require_task_owner(user)
        if (
            not idempotency_key
            or len(idempotency_key) > 128
            or any(
                ord(character) < 33 or ord(character) > 126
                for character in idempotency_key
            )
        ):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "A printable Idempotency-Key of at most 128 characters is required"
                ),
            )
        _executor_for_user(database_path, user)
        metric_catalog = release_store.snapshot() if release_store else None
        template = resolve_template(request.template_id, user.role)
        _validate_report_request(request, user, metric_catalog, template)
        request_payload = {
            "report_request": request.model_dump(mode="json"),
            "metric_catalog": metric_catalog,
            "metric_catalog_sha256": _catalog_snapshot_hash(metric_catalog),
            "report_template": capture_template_snapshot(template),
            "report_template_sha256": _catalog_snapshot_hash(
                capture_template_snapshot(template)
            ),
        }
        try:
            task, created = task_store.enqueue(
                actor_user_id,
                idempotency_key,
                request_payload,
                _task_identity(user),
            )
        except IdempotencyConflict as exc:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "Idempotency-Key was already used for a different report request"
                ),
            ) from exc
        except ReportResultEncryptionError as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report result encryption is not configured",
            ) from exc
        if created:
            response.status_code = status.HTTP_202_ACCEPTED
        task["download_url"] = (
            f"/reports/tasks/{task['task_id']}/export.xlsx"
            if task["status"] == "succeeded"
            else None
        )
        _audit_sink(user)(
            {
                "status": "report_task_queued" if created else "report_task_reused",
                "task_id": task["task_id"],
                "metric_count": len(request.metrics),
            }
        )
        return task

    @app.post("/reports/batches/tasks")
    def create_report_batch_task(
        request: ReportBatchRequest,
        response: Response,
        idempotency_key: Optional[str] = Header(None, alias="Idempotency-Key"),
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if task_store is None or database_path is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report task storage or demo database is not configured",
            )
        actor_user_id = _require_task_owner(user)
        if (
            not idempotency_key
            or len(idempotency_key) > 128
            or any(
                ord(character) < 33 or ord(character) > 126
                for character in idempotency_key
            )
        ):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "A printable Idempotency-Key of at most 128 characters is required"
                ),
            )
        _executor_for_user(database_path, user)
        metric_catalog = release_store.snapshot() if release_store else None
        template_ids = {scenario.report.template_id for scenario in request.scenarios}
        if len(template_ids) != 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="All batch scenarios must use the same report template",
            )
        template_id = next(iter(template_ids))
        template = resolve_template(template_id, user.role)
        for scenario in request.scenarios:
            _validate_report_request(scenario.report, user, metric_catalog, template)
        request_payload = {
            "report_batch": request.model_dump(mode="json"),
            "metric_catalog": metric_catalog,
            "metric_catalog_sha256": _catalog_snapshot_hash(metric_catalog),
            "report_template": capture_template_snapshot(template),
            "report_template_sha256": _catalog_snapshot_hash(
                capture_template_snapshot(template)
            ),
        }
        metric_count = sum(
            len(scenario.report.metrics) for scenario in request.scenarios
        )
        try:
            task, created = task_store.enqueue(
                actor_user_id,
                idempotency_key,
                request_payload,
                _task_identity(user),
            )
        except IdempotencyConflict as exc:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    "Idempotency-Key was already used for a different report request"
                ),
            ) from exc
        except ReportResultEncryptionError as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report result encryption is not configured",
            ) from exc
        if created:
            response.status_code = status.HTTP_202_ACCEPTED
        task["download_url"] = (
            f"/reports/tasks/{task['task_id']}/export.xlsx"
            if task["status"] == "succeeded"
            else None
        )
        _audit_sink(user)(
            {
                "status": "report_batch_task_queued"
                if created
                else "report_task_reused",
                "task_id": task["task_id"],
                "batch_size": len(request.scenarios),
                "metric_count": metric_count,
            }
        )
        return task

    @app.get("/reports/tasks/{task_id}")
    def get_report_task(
        task_id: str,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if task_store is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report task storage is not configured",
            )
        task = task_store.get(task_id, _require_task_owner(user))
        if task is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Task not found"
            )
        task["download_url"] = (
            f"/reports/tasks/{task_id}/export.xlsx"
            if task["status"] == "succeeded"
            else None
        )
        return task

    @app.get("/reports/tasks/{task_id}/export.xlsx")
    def download_report_task(
        task_id: str,
        user: UserRequest = Depends(get_user_from_headers),
    ):
        if task_store is None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Report task storage is not configured",
            )
        actor_user_id = _require_task_owner(user)
        task = task_store.get_report(task_id, actor_user_id)
        if task is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Report not found"
            )
        identity = task["identity"]
        if (
            user.tenant_id != identity.get("tenant_id")
            or user.region_id != identity.get("region_id")
            or user.role != identity.get("role")
            or identity.get("data_source_id") != SCHEMA_METADATA["source"]["id"]
            or identity.get("authorization_policy_version")
            != SCHEMA_METADATA["source"]["policy_version"]
        ):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Report not found"
            )
        request_payload = task_store.get_request(task_id, actor_user_id)
        report_user = UserRequest(
            user_id=identity.get("user_id"),
            tenant_id=identity.get("tenant_id"),
            region_id=identity.get("region_id"),
            role=identity.get("role"),
        )
        result = task["result"]
        template_snapshot = request_payload.get("report_template")
        if template_snapshot is None or _catalog_snapshot_hash(
            template_snapshot
        ) != request_payload.get("report_template_sha256"):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Report not found"
            )
        if "report_batch" in request_payload:
            batch = ReportBatchRequest.model_validate(request_payload["report_batch"])
            template_id = batch.scenarios[0].report.template_id
            template = resolve_template(
                template_id, report_user.role, template_snapshot
            )
            buffer = _render_xlsx_batch_report(
                batch,
                report_user,
                result["generated_at"],
                result["scenarios"],
                template,
            )
        else:
            request = ReportRequest.model_validate(request_payload["report_request"])
            template = resolve_template(
                request.template_id, report_user.role, template_snapshot
            )
            buffer = _render_xlsx_report(
                request,
                report_user,
                result["generated_at"],
                result["definitions"],
                result["result_rows"],
                template,
            )
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": (
                    'attachment; filename="enterprise-analytics-report.xlsx"'
                )
            },
        )

    return app


database_path = os.getenv("ENTERPRISE_TEXT2SQL_DB")
app = create_app(Path(database_path) if database_path else None)
